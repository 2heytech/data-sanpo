import re, urllib.request, io, html
UA={"User-Agent":"Mozilla/5.0"}
def get(u):
    return urllib.request.urlopen(urllib.request.Request(u,headers=UA),timeout=120).read()
base="https://www.e-stat.go.jp/stat-search/files?layout=datalist&toukei=00200521&tstat=000001136464&cycle=0&page={p}"
seen=set()
for tc in ["", "&tclass1=000001136469"]:
  for p in range(1,8):
    try: t=get(base.format(p=p)+tc).decode("utf-8","ignore")
    except Exception as e: print("ERR",p,e); break
    rows=re.findall(r'stat-dataset_list-detail-item-text[^>]*>(.*?)</', t, re.S)
    ids=re.findall(r'statInfId=(\d+)&fileKind=(\d)', t)
    titles=[re.sub(r'\s+',' ',html.unescape(re.sub('<[^>]+>','',x))).strip() for x in re.findall(r'<span class="stat-title-has-data">(.*?)</span>|<a [^>]*class="stat-link_text[^"]*"[^>]*>(.*?)</a>', t, re.S) for x in x if x]
    for blk in re.findall(r'<article class="stat-dataset_list-item.*?</article>', t, re.S):
        tx=re.sub(r'\s+',' ',html.unescape(re.sub('<[^>]+>',' ',blk))).strip()
        sid=re.findall(r'statInfId=(\d+)&fileKind=(\d)', blk)
        key=tx[:80]
        if key in seen: continue
        seen.add(key)
        if re.search(r'補完|昼間|昼夜', tx): print("ROW", tx[:260], sorted(set(sid)))
    if not re.search(r'page=%d'%(p+1), t): break
# 001-1.xlsx の中身
import openpyxl
b=get("https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032214141&fileKind=0")
wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True)
ws=wb[wb.sheetnames[0]]
for i,r in enumerate(ws.iter_rows(values_only=True)):
    if i<14: print("HDR",i,[c for c in r if c is not None][:40])
    vals=[str(c) for c in r if c is not None]
    s=" ".join(vals)
    if ("13000" in s or "東京都" in s) and "0_総数" in s and "00_総数" in s: print("TOKYO",vals[:40]); break
    if i>200000: break
