# 一時的な確認用（元ファイルの形を見る）。PR に入れる前に消す。
import re, urllib.request, io, collections
import openpyxl
UA = {"User-Agent": "Mozilla/5.0"}
def wb(url):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=180) as r:
        return openpyxl.load_workbook(io.BytesIO(r.read()), read_only=True, data_only=True)

print("===== MAFF 2024")
rows = list(wb("https://www.maff.go.jp/j/tokei/kouhyou/sityoson_sansyutu/attach/xls/index-25.xlsx").active.iter_rows(values_only=True))
codes = collections.Counter()
combos = collections.Counter()
for r in rows[1:]:
    combos[(r[6], r[7], r[8])] += 1
    codes[(r[1], r[2])] += 1
print("munis", len(codes))
print("combos(top-level)", [c for c in combos if c[1] == c[2] or c[1].endswith("計") or c[0] == c[1]][:60])
for r in rows[1:]:
    if r[5] in ("千代田区", "横浜市", "特別区部", "八王子市", "大島町") or (isinstance(r[2], int) and r[2] in (0, 100) and r[1] in ("13", "14", 13, 14)) or r[5] in (None, "", "計", "都道府県計"):
        if r[7] in ("農業産出額", "米", "野菜計", "果実計", "畜産計", "野菜", "果実", "畜産"):
            print("  ", r)
vals = collections.Counter(type(r[9]).__name__ if not isinstance(r[9], str) else r[9] for r in rows[1:])
print("value kinds", vals.most_common(10))
print("first codes", sorted(codes)[:20], sorted(codes)[-5:])

for name, url in [("METI2025", "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040480531&fileKind=0"),
                  ("METI2022", "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040137360&fileKind=0"),
                  ("METI2023", "https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040194121&fileKind=0")]:
    print("=====", name)
    w = wb(url)
    ws = w["参考表"]
    rs = list(ws.iter_rows(values_only=True))
    for r in rs[6:10]: print("  H", r[:16])
    tot = [r for r in rs if r[5] == "00"]
    print("  製造業計 rows", len(tot))
    for r in tot[:4] + [r for r in tot if r[4] and re.search(r"札幌市$|横浜市$|千代田区|豊田市|特別区|檜原|青ヶ島|御蔵島", str(r[4]))]:
        print("  ", r[:16])
    kinds = collections.Counter(type(r[13]).__name__ if not isinstance(r[13], str) else r[13] for r in tot)
    print("  shipments kinds", kinds.most_common())
    print("  code lens", collections.Counter(len(str(r[3])) for r in tot))

print("===== COM1 header cols")
rs = list(wb("https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040032863&fileKind=0").active.iter_rows(values_only=True))
for j, (a, b) in enumerate(zip(rs[7], rs[8])):
    if a and (a.startswith("I") or a[:2] in ("56", "57", "58", "59", "60", "61")) and len(a.split("_")[0]) <= 2:
        print("  ", j, a, b, rs[9][j])
print("  levels", collections.Counter(r[0] for r in rs[11:]))
for r in rs[11:]:
    if r[1] and re.search(r"13000|13101|13201|13999|14100|14101|01100|東京都", str(r[1])):
        print("  ", r[:12])
print("===== COM2 levels")
rs = list(wb("https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040032864&fileKind=0").active.iter_rows(values_only=True))
print("  levels", collections.Counter(r[0] for r in rs[11:]))
print("  level2 sample", [r[1] for r in rs[11:] if r[0] == "2"][:5], [r[1] for r in rs[11:] if r[0] == "1"][:30])
